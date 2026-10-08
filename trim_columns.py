"""
Basic Combine -> Basic Combine H & D (Column Trimmer)
------------------------------------------------------
Ei script "Basic combine.xlsx" (shob column shoho, seller export
combine kora file) theke shudhu dorkari 7 ta column rekhe
"Basic combine H & D.xlsx" banay. Baki column gula (VideoUrl, Brand,
Unit, Status, Warranty ityadi) drop kore dey, jate por-porborti
step gula te vul na hoy.

Rakha hobe (ei order e):
    Product ID, Seller Code, Name (English), Highlights(English),
    Highlights(Bengali), Description (English), Description (Bengali)

Header naam e space/capital/parenthesis thakuk ba na thakuk - shob e
automatic match hoy (jemon "Highlights(English)" ba
"Highlights (English)" duitai match korbe).

Kivabe run korbe (cmd/terminal e):
    python trim_columns.py

Tারপর ও tomake 1 ta jinis চাইবে:
    1) "Basic combine.xlsx" file er full path

Output: original file thaka jei folder e ache, oi shei folder e
"Basic combine H & D.xlsx" name e notun file banabe (original file
touch kore na, tai risk nai).
"""

import os
import re
import sys

try:
    from openpyxl import Workbook, load_workbook
except ImportError:
    print("openpyxl install kora nai. Age ei command ta run koro:")
    print("    pip install openpyxl")
    sys.exit(1)


# Ei order onujayi final file e column bosbe. Naam ekhane "canonical"
# form e deya - actual header exact match na hoileo cholbe (normalize
# hoye match hobe).
KEEP_COLUMNS = [
    "Product ID",
    "Seller Code",
    "Name (English)",
    "Highlights(English)",
    "Highlights(Bengali)",
    "Description (English)",
    "Description (Bengali)",
]


def normalize_header(value):
    """Header string ke ekta 'canonical' form e niye ashe, jate
    space thakuk ba na thakuk, capital/small jekono hok - shob e
    match kore. Only a-z0-9 rakha hoy."""
    if value is None:
        return ""
    s = str(value).strip().lower()
    s = re.sub(r"[^a-z0-9]", "", s)
    return s


def trim_file(input_path: str, output_path: str = None) -> str:
    if not os.path.isfile(input_path):
        raise FileNotFoundError(f"File pawa jayni: {input_path}")

    wb = load_workbook(input_path, data_only=True)
    ws = wb.active  # data active sheet e ache dhore nicchi

    headers = [c.value for c in ws[1]]
    norm_to_col = {}
    for i, h in enumerate(headers):
        key = normalize_header(h)
        if key and key not in norm_to_col:  # prothom match ta rakhbo
            norm_to_col[key] = i + 1

    selected_cols = []
    missing = []
    for name in KEEP_COLUMNS:
        col_num = norm_to_col.get(normalize_header(name))
        if col_num:
            selected_cols.append((name, col_num))
        else:
            missing.append(name)

    if not selected_cols:
        raise ValueError("Kono expected column pawa jayni. Header naam check koro.")

    if missing:
        print(f"WARNING: ei column gula source file e pawa jayni, skip kora hocche: {missing}")

    new_wb = Workbook()
    new_ws = new_wb.active
    new_ws.title = ws.title or "Data"

    for j, (name, _) in enumerate(selected_cols, start=1):
        new_ws.cell(row=1, column=j, value=name)

    max_row = ws.max_row
    for row_idx in range(2, max_row + 1):
        for j, (name, src_col) in enumerate(selected_cols, start=1):
            new_ws.cell(row=row_idx, column=j, value=ws.cell(row=row_idx, column=src_col).value)

    # Column width thik kora (rough estimate)
    for j, (name, _) in enumerate(selected_cols, start=1):
        letter = new_ws.cell(row=1, column=j).column_letter
        new_ws.column_dimensions[letter].width = max(18, len(str(name)) + 4)

    if output_path is None:
        output_dir = os.path.dirname(os.path.abspath(input_path))
        output_path = os.path.join(output_dir, "Basic combine H & D.xlsx")

    new_wb.save(output_path)

    print(f"\nDone! {max_row - 1} ta data row, {len(selected_cols)} ta column niye notun file banano hoyeche.")
    print(f"Notun file save hoyeche ekhane:\n    {output_path}\n")
    return output_path


if __name__ == "__main__":
    print("=== Basic Combine -> Basic Combine H & D (Column Trimmer) ===\n")
    input_path = input("'Basic combine.xlsx' file er full path dao: ").strip().strip('"')

    try:
        trim_file(input_path)
    except Exception as e:
        print(f"\nError hoyeche: {e}")
        sys.exit(1)
