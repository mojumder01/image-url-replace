"""Tests with small fake files. Run:  python -m unittest discover tests"""

import os
import shutil
import sys
import tempfile
import unittest
from datetime import datetime
from unittest import mock

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from catprep import config  # noqa: E402
from catprep.common import (StepError, date_from_filename, find_col, is_mapping_file_name,  # noqa: E402
                            norm_header, norm_pid, read_table)
from catprep.run_folder import Run  # noqa: E402
from catprep.steps import STEPS  # noqa: E402
from catprep.steps.s2_export import combine_exports, read_mapping_pids, read_seller_codes  # noqa: E402
from catprep.steps.s4_images import replace_images  # noqa: E402
from catprep.steps.s7_upload import fill_sheet, find_template, split_evenly  # noqa: E402

PROJECT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
BASE = "https://cdn.example.com/product/"
COLS = ["Seller Code", "Product ID", "Name (English)", "Highlights(English)", "Highlights(Bengali)",
        "Description (English)", "Description (Bengali)", "Status"]


def img(url):
    return f'<p>x</p><img class="a" src="{url}">'


def save(path, sheets):
    wb = Workbook()
    wb.remove(wb.active)
    for title, rows in sheets.items():
        ws = wb.create_sheet(title)
        for r in rows:
            ws.append(r)
    wb.save(path)


class Helpers(unittest.TestCase):
    def test_norm_header(self):
        self.assertEqual(norm_header("Highlights (English)"), norm_header("highlights(english)"))

    def test_norm_pid(self):
        for v in (2261883, 2261883.0, "2261883", " 2261883.0 "):
            self.assertEqual(norm_pid(v), "2261883")
        for v in (None, "", float("nan"), "nan"):
            self.assertIsNone(norm_pid(v))

    def test_find_col(self):
        self.assertEqual(find_col(["a", "*Product Id"], "Product ID", "*Product Id"), 1)
        self.assertIsNone(find_col(["a"], "b"))

    def test_mapping_file_name(self):
        self.assertTrue(is_mapping_file_name("image_links_mapping (28.09.2026) Bakaul.xlsx"))
        self.assertTrue(is_mapping_file_name("Image Links Bakaul 01.10.2026.xlsx"))
        self.assertFalse(is_mapping_file_name("image_http_url_add.xlsx"))
        self.assertFalse(is_mapping_file_name("Basic combine.xlsx"))

    def test_date_from_filename(self):
        self.assertEqual(date_from_filename("image_links_mapping (28.09.2026) Bakaul.xlsx"), datetime(2026, 9, 28))
        self.assertEqual(date_from_filename("x_28-09-2026.xlsx"), datetime(2026, 9, 28))
        self.assertEqual(date_from_filename("x 2026-09-28.xlsx"), datetime(2026, 9, 28))
        self.assertIsNone(date_from_filename("no date.xlsx"))

    def test_replace_images(self):
        html = '<img src="a"><p>t</p><img alt="" src=\'b\'><img src="c">'
        new, replaced, total = replace_images(html, ["N1", "N2"])
        self.assertEqual((replaced, total), (2, 3))
        self.assertEqual(new, '<img src="N1"><p>t</p><img alt="" src=\'N2\'><img src="c">')

    def test_split_evenly(self):
        self.assertEqual([len(c) for c in split_evenly(list(range(10)), 7500, 2)], [5, 5])
        self.assertEqual([len(c) for c in split_evenly(list(range(16000)), 7500, 2)], [5334, 5333, 5333])
        self.assertEqual([len(c) for c in split_evenly([1], 7500, 2)], [1])

    def test_config_defaults_and_paths(self):
        cfg = config.load(PROJECT)
        self.assertEqual(cfg["claude"]["rows_per_part"], 3500)
        self.assertTrue(cfg.path(r"templates\x.xlsx").endswith(os.path.join("templates", "x.xlsx")))


class RunFolderTest(unittest.TestCase):
    def test_name_and_status(self):
        tmp = tempfile.mkdtemp()
        try:
            r1 = Run.create(tmp, datetime(2026, 9, 28), now=datetime(2026, 9, 29, 15, 37))
            self.assertEqual(r1.name, "September 2026_28-09-2026_15-37")
            r2 = Run.create(tmp, datetime(2026, 9, 28), now=datetime(2026, 9, 29, 15, 37))
            self.assertEqual(r2.name, "September 2026_28-09-2026_15-37-2")
            r1.step_done("urls", cells_updated=3)
            self.assertTrue(r1.is_done("urls"))
            self.assertFalse(r1.is_done("trim"))
            self.assertEqual(len(Run.all_runs(tmp)), 2)
        finally:
            shutil.rmtree(tmp)


class PipelineTest(unittest.TestCase):
    """Runs every step except the browser download on tiny fake data."""

    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        self.cfg = config.load(PROJECT)
        self.cfg["output_root"] = os.path.join(self.tmp, "out")
        self.cfg["base_url"] = BASE
        self.cfg["claude"]["rows_per_part"] = 2
        self.cfg["claude"]["fallback_done_folder"] = os.path.join(self.tmp, "nothing")

        mapping = os.path.join(self.tmp, "image_links_mapping (28.09.2026) Bakaul.xlsx")
        save(mapping, {
            "Data Mapping": [["Product ID", "image1", "image2"],
                             ["101", "k1.webp", "k2.webp"],
                             ["102", "k3.webp", None],
                             ["103", "https://already/url.webp", None],
                             ["104", "k5.webp", None],
                             ["105", "k6.webp", None]],
            "Seller Info": [["Seller ID", "Seller Code"], ["1", "SLC1"], ["2", "SLC2"]],
        })
        self.run_ = Run.create(self.cfg["output_root"], datetime(2026, 9, 28))
        shutil.copy2(mapping, self.run_.folder("input"))
        self.run_.set(mapping_file=os.path.basename(mapping))

        # Fake seller export: 2 header rows, Product Update sheet, a dropdown.
        downloads = os.path.join(self.run_.folder("export"), "downloads")
        os.makedirs(downloads)
        wb = Workbook()
        ws = wb.active
        ws.title = "Product Update"
        ws.append(["Basic"] * 3 + ["Description"] * 4 + ["Other"])
        ws.append(COLS)
        ws.append(["SLC1", 101, "P101", img("old1"), None, img("old1") + img("old2"), None, "Active"])
        ws.append(["SLC1", 102, "P102", img("old"), img("old"), "<p>d</p>", None, "Active"])
        ws.append(["SLC1", 103, "P103", None, None, img("a") + img("b"), None, "Active"])  # too few
        ws.append(["SLC2", 104, "P104", "<ul><li>h</li></ul>", None, "<p>d</p>", None, "Active"])
        ws.append(["SLC2", 105, "P105", "<ul><li>h</li></ul>", None, "<p>d</p>", None, "Active"])
        ws.append(["SLC2", 999, "Not in mapping", None, None, None, None, "Active"])
        dv = DataValidation(type="list", formula1='"Active,Inactive"')
        dv.add("H3:H8")
        ws.add_data_validation(dv)
        wb.create_sheet("DropdownData").append(["Active"])
        wb.save(os.path.join(downloads, "batch1_export.xlsx"))

        class Ctx:
            pass
        self.ctx = Ctx()
        self.ctx.run, self.ctx.cfg, self.ctx.options = self.run_, self.cfg, {}

    def tearDown(self):
        shutil.rmtree(self.tmp)

    def step(self, name):
        self.run_.step_done(name, **(STEPS[name][0].run(self.ctx) or {}))

    def prepare(self):
        self.step("urls")
        mapping = self.run_.mapping_file()
        self.assertEqual(read_seller_codes(mapping, "Seller Info"), ["SLC1", "SLC2"])
        files = [os.path.join(self.run_.folder("export"), "downloads", "batch1_export.xlsx")]
        total, kept = combine_exports(files, self.run_.file("combine"), read_mapping_pids(mapping, "Data Mapping"))
        self.assertEqual((total, kept), (6, 5))
        for name in ("trim", "images", "split"):
            self.step(name)

    def claude(self, drop=()):
        """Fake Claude: mark the HTML columns as cleaned, optionally drop products."""
        for i, part in enumerate(sorted(os.listdir(self.run_.folder("for_claude")))):
            if not part.startswith("Part-"):
                continue
            h, rows = read_table(os.path.join(self.run_.folder("for_claude"), part))
            rows = [r[:3] + [f"CLEAN-{r[0]}-{j}" for j in range(4)] for r in rows if r[0] not in drop]
            h = [x.replace("(", " (") for x in h]  # Claude may change header spacing
            save(os.path.join(self.run_.folder("claude_done"), f"out{i}.xlsx"), {"Sheet1": [h] + rows})

    def test_full_pipeline(self):
        self.prepare()

        urls = read_table(self.run_.file("urls"), sheet="Data Mapping")[1]
        self.assertEqual(urls[0][1:], [BASE + "k1.webp", BASE + "k2.webp"])
        self.assertEqual(urls[2][1], "https://already/url.webp")

        h, rows = read_table(self.run_.file("trim"))
        self.assertEqual(h, self.cfg["columns"]["keep"])
        self.assertEqual(len(rows), 5)

        h, good = read_table(self.run_.file("images"), sheet="Data")
        _, review = read_table(self.run_.file("images"), sheet="Needs Review")
        self.assertEqual([r[0] for r in good], ["101", "102", "104", "105"])
        self.assertEqual([(r[0], r[-1]) for r in review], [("103", "Insufficient Mapped Images")])
        p101 = good[0]
        self.assertIn(BASE + "k1.webp", p101[3])
        self.assertIn(BASE + "k2.webp", p101[5])
        self.assertNotIn("old", p101[3] + p101[5])

        parts = sorted(f for f in os.listdir(self.run_.folder("for_claude")) if f.startswith("Part-"))
        self.assertEqual(parts, ["Part-1.xlsx", "Part-2.xlsx"])
        self.assertTrue(os.path.isfile(os.path.join(self.run_.folder("for_claude"), "H&D_improved_prompt.md")))

        self.claude()
        self.step("merge")
        self.step("upload")

        out = self.run_.folder("upload")
        up_rows = []
        for f in ("update_part1.xlsx", "update_part2.xlsx"):
            wb = load_workbook(os.path.join(out, f))
            self.assertIn("DropdownData", wb.sheetnames)
            ws = wb["Product Update"]
            self.assertEqual(ws.cell(row=1, column=1).value, "Basic")
            up_rows += [list(r) for r in ws.iter_rows(min_row=3, values_only=True)]
        self.assertEqual(sorted(norm_pid(r[1]) for r in up_rows), ["101", "102", "104", "105"])
        row101 = next(r for r in up_rows if norm_pid(r[1]) == "101")
        self.assertEqual(row101[3:7], ["CLEAN-101-0", "CLEAN-101-1", "CLEAN-101-2", "CLEAN-101-3"])
        self.assertEqual(row101[7], "Active")

        _, appr = read_table(self.run_.file("approval"))
        self.assertEqual(sorted(norm_pid(r[1]) for r in appr), ["101", "102", "104", "105"])
        self.assertTrue(all(str(r[2]) == "1" for r in appr))
        _, left = read_table(self.run_.file("not_uploaded"))
        self.assertEqual([(r[0], r[3]) for r in left], [("103", "Needs Review: Insufficient Mapped Images")])

    def test_missing_rows_stop_upload(self):
        self.prepare()
        self.claude(drop={"104"})
        self.step("merge")
        with mock.patch("catprep.steps.s7_upload.interactive", return_value=False):
            with self.assertRaises(StepError):
                STEPS["upload"][0].run(self.ctx)
        self.ctx.options["allow_missing"] = True
        self.step("upload")
        _, left = read_table(self.run_.file("not_uploaded"))
        self.assertIn(("104", "Missing from Claude output"), [(r[0], r[3]) for r in left])

    def test_needs_review_never_uploaded(self):
        """Even if Claude's files contain a Needs Review product, it is not uploaded."""
        self.prepare()
        self.claude()
        extra = os.path.join(self.run_.folder("claude_done"), "extra.xlsx")
        save(extra, {"Sheet1": [self.cfg["columns"]["keep"], ["103", "SLC1", "P103", "a", "b", "c", "d"]]})
        self.step("merge")
        self.step("upload")
        _, appr = read_table(self.run_.file("approval"))
        self.assertNotIn("103", [norm_pid(r[1]) for r in appr])
        _, left = read_table(self.run_.file("not_uploaded"))
        self.assertEqual([r[0] for r in left], ["103"])

    def test_dropdowns_extend_past_template_rows(self):
        """Writing more rows than the template has keeps formats and stretches dropdowns."""
        template = os.path.join(self.run_.folder("export"), "downloads", "batch1_export.xlsx")
        self.assertEqual(find_template(os.path.dirname(template)), template)
        wb = load_workbook(template)
        ws = wb["Product Update"]
        fill_sheet(ws, [["S", str(i)] + [None] * 5 + ["Active"] for i in range(20)], first_data_row=3)
        ranges = [str(r) for dv in ws.data_validations.dataValidation for r in dv.sqref.ranges]
        self.assertEqual(ranges, ["H3:H22"])
        self.assertEqual(ws.cell(row=22, column=2).value, "19")
        self.assertEqual(ws.cell(row=1, column=1).value, "Basic")

if __name__ == "__main__":
    unittest.main()
