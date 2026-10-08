"""Step 2: export every seller's products from seller-admin and combine them.

seller-admin.cartup.com -> Products -> Multi Seller Bulk Update -> Export Excel File
-> seller codes (in batches) -> Download -> Export History -> wait for "completed"
-> download. Then all downloads are joined into 'Basic combine.xlsx':
  - sheet 'Combined': every row
  - sheet 'Workings': only Product IDs that are in the mapping file (active sheet)

Login: see catprep/site.py (saved session, credentials.json, OTP from email or typed).
Batches already downloaded (batchN_*.xlsx) are skipped, so a failed run can
simply be started again.

Manual mode (--manual-export): download the export files yourself, put them in
the run's 2_export/downloads folder, press Enter; they are combined the same way.
"""

import os
import time

from openpyxl import Workbook, load_workbook

from .. import site
from ..common import (PID_NAMES, StepError, find_col, info, interactive, norm_header, norm_pid,
                      pick_sheet, warn, write_table)
from ..site import click, wait_overlay_clear


# ---------------------------------------------------------------------------
# Reading the mapping file
# ---------------------------------------------------------------------------
def read_seller_codes(mapping_path, sheet):
    wb = load_workbook(mapping_path, read_only=True, data_only=True)
    try:
        if norm_header(sheet) not in [norm_header(s) for s in wb.sheetnames]:
            raise StepError(f"Sheet '{sheet}' not found in the mapping file. Sheets: {wb.sheetnames}")
        rows = list(pick_sheet(wb, sheet).iter_rows(values_only=True))
    finally:
        wb.close()
    col = find_col(rows[0] if rows else [], "Seller Code")
    if col is None:
        raise StepError(f"No 'Seller Code' column in sheet '{sheet}'.")
    codes = []
    for row in rows[1:]:
        code = str(row[col]).strip() if col < len(row) and row[col] is not None else ""
        if code and code not in codes:
            codes.append(code)
    return codes


def read_mapping_pids(mapping_path, sheet):
    wb = load_workbook(mapping_path, read_only=True, data_only=True)
    try:
        rows = list(pick_sheet(wb, sheet).iter_rows(values_only=True))
    finally:
        wb.close()
    col = find_col(rows[0] if rows else [], *PID_NAMES)
    if col is None:
        raise StepError(f"No 'Product ID' column in mapping sheet '{sheet}'.")
    return {p for p in (norm_pid(r[col]) for r in rows[1:] if col < len(r)) if p}


# ---------------------------------------------------------------------------
# Combining downloaded export files
# ---------------------------------------------------------------------------
def combine_exports(files, out_path, valid_pids):
    """Export files have 2 header rows: row 1 = groups (dropped), row 2 = columns."""
    wb = Workbook()
    combined = wb.active
    combined.title = "Combined"
    header = None
    all_rows = []
    for path in files:
        try:
            src = load_workbook(path, read_only=True, data_only=True)
            rows = list(src.active.iter_rows(values_only=True))
            src.close()
        except Exception as e:
            warn(f"Cannot read {os.path.basename(path)}, skipped: {e}")
            continue
        if len(rows) < 2:
            warn(f"{os.path.basename(path)} has no column header row, skipped.")
            continue
        if header is None:
            header = list(rows[1])
            combined.append(header)
        for row in rows[2:]:
            if any(v is not None for v in row):
                combined.append(list(row))
                all_rows.append(list(row))
    if header is None:
        raise StepError("No usable export file to combine.")

    pid_col = find_col(header, *PID_NAMES)
    if pid_col is None:
        raise StepError("Export files have no 'Product ID' column.")
    workings = wb.create_sheet("Workings")
    workings.append(header)
    kept = 0
    for row in all_rows:
        if norm_pid(row[pid_col]) in valid_pids:
            workings.append(row)
            kept += 1
    wb.active = wb.sheetnames.index("Workings")
    wb.save(out_path)
    info(f"Combined {len(files)} file(s): {len(all_rows)} rows, {kept} match the mapping file -> {out_path}")
    return len(all_rows), kept


# ---------------------------------------------------------------------------
# Browser robot (Playwright, PrimeReact UI)
# ---------------------------------------------------------------------------
def _open_export_page(page):
    click(page.locator("text=Products").first, page, "Products menu")
    page.wait_for_timeout(400)
    click(page.locator("text=Multi Seller Bulk").first, page, "Multi Seller Bulk Update")
    page.wait_for_load_state("domcontentloaded")
    page.wait_for_timeout(500)
    tab = page.locator("text=Export Excel File").first
    if tab.count():
        click(tab, page, "Export Excel File tab", 3000)
        page.wait_for_timeout(300)


def _submit_batch(page, codes):
    field = page.get_by_placeholder("Enter Seller Code(s)")
    field.fill("")
    field.fill(",".join(codes))
    click(page.locator("text=All").first, page, "Status: All", 2000)
    click(page.locator("text=Basic Information").first, page, "Edit: Basic Information", 2000)
    click(page.locator("button:has-text('Download')").first, page, "Download (submit) button", 3000)


def _filter_history(page, name):
    if not name:
        return
    field = page.get_by_placeholder("Downloaded By")
    if field.count():
        field.fill("")
        field.fill(name)
        click(page.locator("button:has-text('Search')").first, page, "Search", 2500)
        page.wait_for_timeout(400)


def _wait_and_download(page, out_dir, label, poll_s, timeout_s):
    """Wait for the newest Export History row to be 'completed', then download it."""
    start = time.time()

    def refresh():
        btn = page.locator("button:has-text('Refresh')").first
        if btn.count():
            click(btn, page, "Refresh", 3000)

    while time.time() - start < timeout_s:
        wait_overlay_clear(page)
        row = page.locator("table tbody tr").first
        text = ""
        if row.count():
            try:
                text = row.inner_text(timeout=1500).lower()
            except Exception:
                pass
        if "failed" in text:
            return "failed", ""
        if "completed" in text:
            link = row.locator("a:has-text('Download')").first
            if link.count():
                try:
                    with page.expect_download(timeout=120000) as d:
                        click(link, page, "row Download link", 2500)
                    target = os.path.join(out_dir, f"{label}_{d.value.suggested_filename or 'export.xlsx'}")
                    d.value.save_as(target)
                    info(f"Downloaded {os.path.basename(target)}")
                    return "completed", target
                except Exception as e:
                    warn(f"Download failed, will retry: {e}")
        info(f"  not ready yet, checking again in {poll_s}s ...")
        refresh()
        page.wait_for_timeout(poll_s * 1000)
    return "timeout", ""


def download_exports(cfg, codes, downloads_dir):
    ex = cfg["export"]
    size = int(ex["batch_size"])
    batches = [codes[i:i + size] for i in range(0, len(codes), size)]
    info(f"{len(codes)} seller codes -> {len(batches)} batches of up to {size}.")
    results = []
    with site.session(cfg) as page:
        name = site.display_name(page, cfg)
        _open_export_page(page)
        for i, batch in enumerate(batches, start=1):
            label = f"batch{i}"
            done = [f for f in os.listdir(downloads_dir) if f.startswith(label + "_")]
            if done:
                info(f"Batch {i}/{len(batches)} already downloaded ({done[0]}) - skipped.")
                results.append([i, ",".join(batch), "already downloaded", done[0]])
                continue
            info(f"\nBatch {i}/{len(batches)}: {','.join(batch)}")
            _submit_batch(page, batch)
            page.wait_for_timeout(1500)
            _filter_history(page, name)
            state, path = _wait_and_download(page, downloads_dir, label,
                                             int(ex["poll_seconds"]), int(ex["timeout_seconds"]))
            if state != "completed":
                warn(f"Batch {i}: export {state}.")
            results.append([i, ",".join(batch), state, os.path.basename(path)])
    return results


# ---------------------------------------------------------------------------
def run(ctx):
    cfg, run_ = ctx.cfg, ctx.run
    mapping = run_.mapping_file()
    downloads = os.path.join(run_.folder("export"), "downloads")
    os.makedirs(downloads, exist_ok=True)

    if ctx.options.get("manual_export"):
        info("MANUAL EXPORT: download the seller export files from seller-admin and put them here:")
        info(f"    {downloads}")
        if interactive():
            input("Press ENTER when the files are in the folder ... ")
    else:
        codes = read_seller_codes(mapping, cfg["mapping"]["seller_sheet"])
        if not codes:
            raise StepError("No seller codes in the mapping file.")
        results = download_exports(cfg, codes, downloads)
        write_table(run_.file("export_report"), ["Batch", "Seller Codes", "Status", "File"], results,
                    sheet="Batches")
        bad = [r for r in results if r[2] not in ("completed", "already downloaded")]
        if bad:
            raise StepError(f"{len(bad)} batch(es) did not download: {[r[0] for r in bad]}. "
                            "Run the same command again with --resume to retry only those.")

    files = sorted(os.path.join(downloads, f) for f in os.listdir(downloads)
                   if f.lower().endswith(".xlsx") and not f.startswith("~$"))
    if not files:
        raise StepError(f"No export files in {downloads}")
    valid = read_mapping_pids(mapping, cfg["mapping"]["data_sheet"])
    total, kept = combine_exports(files, run_.file("combine"), valid)
    if kept == 0:
        raise StepError("None of the exported products are in the mapping file - check the seller codes.")
    return {"export_files": len(files), "rows_exported": total, "rows_in_mapping": kept}
