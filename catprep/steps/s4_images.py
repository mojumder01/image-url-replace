"""Step 4: put the new image URLs into the HTML columns.

For each product, in each HTML column: the 1st <img src> gets image1, the 2nd
gets image2, ... (every column starts again at image1). Extra <img> tags with no
new image keep their old URL.

Rows that cannot be fully fixed go to the 'Needs Review' sheet (with an Issue):
  - No Mapping Found            Product ID is not in the mapping file
  - Insufficient Mapped Images  more <img> tags than new images
They are NOT sent to Claude and NOT uploaded.
"""

import os
import re

from openpyxl import load_workbook

from ..common import (PID_NAMES, StepError, add_report_sheet, add_table_sheet, find_col, info,
                      new_workbook, norm_header, norm_pid, read_table, timestamp)

IMG_SRC = re.compile(r'(<img[^>]*\ssrc=)(["\'])(.*?)\2', re.IGNORECASE)

NO_MAPPING = "No Mapping Found"
TOO_FEW_IMAGES = "Insufficient Mapped Images"


def load_image_map(urls_path):
    """{product_id: [url1, url2, ...]} from the first sheet with a Product ID column."""
    wb = load_workbook(urls_path, read_only=True, data_only=True)
    try:
        for ws in wb.worksheets:
            if norm_header(ws.title) == "report":
                continue
            rows = ws.iter_rows(values_only=True)
            headers = list(next(rows, []))
            pid_col = find_col(headers, *PID_NAMES)
            if pid_col is None:
                continue
            img_cols = [i for i, h in enumerate(headers) if h and "image" in norm_header(h)]
            result = {}
            for row in rows:
                pid = norm_pid(row[pid_col]) if pid_col < len(row) else None
                if not pid:
                    continue
                urls = [str(row[c]).strip() for c in img_cols if c < len(row) and row[c]]
                if urls:
                    result[pid] = urls
            return result
    finally:
        wb.close()
    raise StepError("No sheet with a 'Product ID' column in the image URL file.")


def replace_images(html, urls):
    """Returns (new_html, replaced, total_img_tags)."""
    state = {"replaced": 0, "total": 0}

    def sub(m):
        state["total"] += 1
        i = state["replaced"]
        if i < len(urls):
            state["replaced"] += 1
            return f"{m.group(1)}{m.group(2)}{urls[i]}{m.group(2)}"
        return m.group(0)

    return IMG_SRC.sub(sub, html), state["replaced"], state["total"]


def merge(trim_path, urls_path, html_columns, out_path):
    image_map = load_image_map(urls_path)
    headers, rows = read_table(trim_path)
    pid_col = find_col(headers, *PID_NAMES)
    if pid_col is None:
        raise StepError("No 'Product ID' column in the trimmed file.")
    html_cols = [find_col(headers, name) for name in html_columns]
    html_cols = [c for c in html_cols if c is not None]

    good, review = [], []
    stats = {"no_mapping": 0, "too_few": 0, "tags_found": 0, "tags_replaced": 0}
    for row in rows:
        pid = norm_pid(row[pid_col])
        if pid is None:
            good.append(row)
            continue
        urls = image_map.get(pid)
        if not urls:
            stats["no_mapping"] += 1
            review.append(row + [NO_MAPPING])
            continue
        short = False
        for c in html_cols:
            if not row[c]:
                continue
            new_html, replaced, total = replace_images(str(row[c]), urls)
            stats["tags_found"] += total
            stats["tags_replaced"] += replaced
            short = short or replaced < total
            row[c] = new_html
        if short:
            stats["too_few"] += 1
            review.append(row + [TOO_FEW_IMAGES])
        else:
            good.append(row)

    wb = new_workbook()
    add_table_sheet(wb, "Data", headers, good, text_cols=[pid_col])
    add_table_sheet(wb, "Needs Review", headers + ["Issue"], review, text_cols=[pid_col],
                    header_fill="ED0A5E")
    add_report_sheet(wb, "Merge Report", "Image Merge into HTML - Summary", [
        ("Main file:", os.path.basename(trim_path)),
        ("Image URL file:", os.path.basename(urls_path)),
        ("Run time:", timestamp()),
        ("Product rows processed:", len(rows)),
        ("Rows ready for Claude ('Data' sheet):", len(good)),
        ("Rows moved to 'Needs Review':", len(review)),
        ("  - No mapping found:", stats["no_mapping"]),
        ("  - Insufficient mapped images:", stats["too_few"]),
        ("<img> tags found:", stats["tags_found"]),
        ("<img> tags replaced:", stats["tags_replaced"]),
        ("HTML columns processed:", ", ".join(headers[c] for c in html_cols) or "(none found!)"),
    ])
    wb.save(out_path)
    return len(good), len(review), stats


def run(ctx):
    out = ctx.run.file("images")
    good, review, stats = merge(ctx.run.file("trim"), ctx.run.file("urls"),
                                ctx.cfg["columns"]["html"], out)
    info(f"Ready for Claude: {good} rows | Needs Review: {review} "
         f"(no mapping {stats['no_mapping']}, too few images {stats['too_few']}) -> {out}")
    return {"rows_ok": good, "needs_review": review, "tags_replaced": stats["tags_replaced"]}
