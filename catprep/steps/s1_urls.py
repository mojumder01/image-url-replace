"""Step 1: turn image keys in the mapping file into full URLs.

'dcf2879f-....webp'  ->  'https://sl-dev-s3.s3.amazonaws.com/product/dcf2879f-....webp'

Every column whose header contains "image" is changed, in every sheet.
Cells that already start with http(s):// and empty cells are left alone.
"""

import os

from openpyxl import load_workbook

from ..common import StepError, add_report_sheet, info, timestamp


def add_urls(mapping_path, base_url, out_path):
    """Returns a dict of counts. Writes out_path with an extra 'Report' sheet."""
    if not base_url:
        raise StepError("base_url is empty in config.toml.")
    if not base_url.endswith("/"):
        base_url += "/"

    wb = load_workbook(mapping_path)
    totals = {"updated": 0, "already_url": 0, "blank": 0}
    per_sheet = []
    for ws in wb.worksheets:
        if ws.max_row < 2:
            continue
        image_cols = [c for c in range(1, ws.max_column + 1)
                      if "image" in str(ws.cell(row=1, column=c).value or "").lower()]
        if not image_cols:
            continue
        counts = {"updated": 0, "already_url": 0, "blank": 0}
        for r in range(2, ws.max_row + 1):
            for c in image_cols:
                cell = ws.cell(row=r, column=c)
                value = str(cell.value).strip() if cell.value is not None else ""
                if not value:
                    counts["blank"] += 1
                elif value.lower().startswith(("http://", "https://")):
                    counts["already_url"] += 1
                else:
                    cell.value = base_url + value
                    counts["updated"] += 1
        for k in totals:
            totals[k] += counts[k]
        per_sheet.append((ws.title, len(image_cols), counts))

    if "Report" in wb.sheetnames:
        del wb["Report"]
    lines = [
        ("Source file:", os.path.basename(mapping_path)),
        ("Base URL used:", base_url),
        ("Run time:", timestamp()),
        ("Cells updated:", totals["updated"]),
        ("Already a URL (skipped):", totals["already_url"]),
        ("Blank cells (skipped):", totals["blank"]),
    ]
    for title, n_cols, c in per_sheet:
        lines.append((f"Sheet '{title}':", f"{n_cols} image columns, {c['updated']} updated"))
    add_report_sheet(wb, "Report", "Image URL Add - Summary", lines)
    wb.save(out_path)
    return totals


def run(ctx):
    out = ctx.run.file("urls")
    totals = add_urls(ctx.run.mapping_file(), ctx.cfg["base_url"], out)
    if totals["updated"] == 0 and totals["already_url"] == 0:
        raise StepError("No image cells found in the mapping file (no column with 'image' in its header).")
    info(f"{totals['updated']} image keys turned into URLs -> {out}")
    return {"cells_updated": totals["updated"]}
