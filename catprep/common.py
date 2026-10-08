"""Shared helpers used by every step: name matching, Excel read/write, console."""

import os
import re
import sys
from datetime import datetime

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter


# ---------------------------------------------------------------------------
# Matching headers and Product IDs
# ---------------------------------------------------------------------------
def norm_header(value):
    """'Highlights (English)', 'highlights(english)' -> 'highlightsenglish'."""
    if value is None:
        return ""
    return re.sub(r"[^a-z0-9]", "", str(value).strip().lower())


def norm_pid(value):
    """2261883 / 2261883.0 / ' 2261883 ' -> '2261883'. Blank -> None."""
    if value is None:
        return None
    if isinstance(value, float):
        if value != value:  # NaN
            return None
        if value.is_integer():
            return str(int(value))
    s = str(value).strip()
    if s.endswith(".0"):
        s = s[:-2]
    if not s or s.lower() == "nan":
        return None
    return s


def find_col(headers, *names):
    """0-based index of the first header matching any of names, or None."""
    normalized = [norm_header(h) for h in headers]
    for name in names:
        key = norm_header(name)
        if key in normalized:
            return normalized.index(key)
    return None


PID_NAMES = ("Product ID", "ProductID", "Product Id", "*Product Id")


# ---------------------------------------------------------------------------
# Finding the mapping file and its date
# ---------------------------------------------------------------------------
def is_mapping_file_name(name):
    """The name must contain 'image', 'link' and one of mapping/maping/bakaul.
    Example: 'image_links_mapping (28.09.2026) Bakaul.xlsx'."""
    low = name.lower()
    tokens = set(re.split(r"[^a-z0-9]+", low))
    has_kind = bool(tokens & {"mapping", "maping", "bakaul"})
    return has_kind and "image" in low and "link" in low


def find_mapping_files(folders):
    """All mapping-like .xlsx files in folders, newest first."""
    found = []
    for folder in folders:
        if not os.path.isdir(folder):
            continue
        for name in os.listdir(folder):
            if name.lower().endswith(".xlsx") and not name.startswith("~$") and is_mapping_file_name(name):
                full = os.path.join(folder, name)
                found.append((os.path.getmtime(full), full))
    found.sort(reverse=True)
    return [full for _, full in found]


def date_from_filename(name):
    """Find DD.MM.YYYY, DD-MM-YYYY or YYYY-MM-DD in a file name."""
    base = os.path.basename(name or "")
    for pattern, order in ((r"(\d{2})[._-](\d{2})[._-](\d{4})", "dmy"),
                           (r"(\d{4})[._-](\d{2})[._-](\d{2})", "ymd")):
        for m in re.finditer(pattern, base):
            a, b, c = (int(x) for x in m.groups())
            d, mo, y = (a, b, c) if order == "dmy" else (c, b, a)
            try:
                return datetime(y, mo, d)
            except ValueError:
                continue
    return None


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------
def list_xlsx(folder):
    """Sorted .xlsx files in a folder (Excel lock files skipped)."""
    if not os.path.isdir(folder):
        return []
    return [os.path.join(folder, f) for f in sorted(os.listdir(folder), key=natural_key)
            if f.lower().endswith(".xlsx") and not f.startswith("~$")]


def natural_key(text):
    """Sort 'Part-2' before 'Part-10'."""
    return [int(t) if t.isdigit() else t.lower() for t in re.split(r"(\d+)", text)]


def pick_sheet(wb, preferred=None):
    """Sheet by (loosely matched) name, else the active sheet."""
    if preferred:
        for ws in wb.worksheets:
            if norm_header(ws.title) == norm_header(preferred):
                return ws
    return wb.active


def read_table(path, sheet=None, header_row=1):
    """Return (headers, rows) where rows are lists. Fully empty rows are dropped."""
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = pick_sheet(wb, sheet)
        headers, rows = [], []
        for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
            if i < header_row:
                continue
            if i == header_row:
                headers = list(row)
                continue
            if any(v is not None and str(v).strip() != "" for v in row):
                rows.append(list(row))
    finally:
        wb.close()
    while headers and headers[-1] is None:
        headers.pop()
    width = len(headers)
    rows = [(r + [None] * width)[:width] for r in rows]
    return headers, rows


def add_table_sheet(wb, title, headers, rows, text_cols=(), header_fill=None):
    """Add a sheet with a header row and data. text_cols are stored as text."""
    from openpyxl.styles import Font, PatternFill

    ws = wb.create_sheet(title)
    ws.append(list(headers))
    for row in rows:
        ws.append(list(row))
    for c in text_cols:
        for r in range(2, ws.max_row + 1):
            cell = ws.cell(row=r, column=c + 1)
            if cell.value is not None:
                cell.value = str(cell.value)
                cell.number_format = "@"
    bold = Font(bold=True, color="FFFFFF" if header_fill else None)
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=c)
        cell.font = bold
        if header_fill:
            cell.fill = PatternFill(start_color=header_fill, end_color=header_fill, fill_type="solid")
        ws.column_dimensions[get_column_letter(c)].width = max(14, min(40, len(str(headers[c - 1] or "")) + 4))
    ws.freeze_panes = "A2"
    return ws


def add_report_sheet(wb, title, heading, lines):
    """Simple two-column summary sheet: [(label, value), ...]."""
    from openpyxl.styles import Font

    ws = wb.create_sheet(title)
    ws["A1"] = heading
    ws["A1"].font = Font(bold=True, size=14)
    for i, (label, value) in enumerate(lines, start=3):
        ws.cell(row=i, column=1, value=label).font = Font(bold=True)
        ws.cell(row=i, column=2, value=value)
    ws.column_dimensions["A"].width = 55
    ws.column_dimensions["B"].width = 60
    return ws


def new_workbook():
    """Empty workbook (the default sheet is removed)."""
    wb = Workbook()
    wb.remove(wb.active)
    return wb


def write_table(path, headers, rows, sheet="Data", text_cols=()):
    wb = new_workbook()
    add_table_sheet(wb, sheet, headers, rows, text_cols=text_cols)
    wb.save(path)
    return path


def timestamp():
    return datetime.now().strftime("%Y-%m-%d %H:%M:%S")


# ---------------------------------------------------------------------------
# Console
# ---------------------------------------------------------------------------
class StepError(Exception):
    """A step cannot continue. The message is shown to the user as-is."""


def info(msg=""):
    print(msg, flush=True)


def warn(msg):
    print(f"WARNING: {msg}", flush=True)


def heading(title):
    info("")
    info("=" * 64)
    info(f"  {title}")
    info("=" * 64)


def interactive():
    return sys.stdin is not None and sys.stdin.isatty()


def ask_yes_no(question, default=False):
    """Ask in a terminal; without one (scheduled/piped run) use default."""
    if not interactive():
        info(f"{question} -> {'yes' if default else 'no'} (no terminal, using default)")
        return default
    suffix = " [Y/n]: " if default else " [y/N]: "
    try:
        answer = input(question + suffix).strip().lower()
    except EOFError:
        return default
    return default if not answer else answer in ("y", "yes")
