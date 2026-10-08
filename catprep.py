"""Catalogue Prep v2

Fully automatic:
  python catprep.py auto               new run, everything: urls -> export -> trim -> images
                                       -> clean -> upload files -> upload to seller-admin
  python catprep.py auto --if-new      same, but only when Data has a mapping file not done yet
                                       (used by the Windows scheduled task)
  python catprep.py auto --resume      continue the newest run from its first unfinished step
  python catprep.py schedule --watch   check Data every few minutes (Windows Task Scheduler)
  python catprep.py schedule --daily   run once a day at a fixed time
  python catprep.py schedule --remove  remove the scheduled tasks

Manual Claude route:
  python catprep.py prepare [--resume] steps up to the Part files for Claude
  python catprep.py finish             merge Claude's files, build + upload files

Other:
  python catprep.py step <name>        re-run one step on the newest run
  python catprep.py status | runs      progress of the newest run | list all runs
  python catprep.py check <file> [file2]   look inside an Excel file (read-only)

Add --run "<run folder>" to work on an older run instead of the newest.
Settings: config.toml
"""

import argparse
import os
import shutil
import subprocess
import sys
import time
import traceback
from datetime import datetime

from catprep import config as config_mod
from catprep import notify
from catprep.common import (PID_NAMES, StepError, date_from_filename, find_col, find_mapping_files,
                            heading, info, norm_pid, warn)
from catprep.run_folder import Run
from catprep.steps import FINISH, PREPARE, STEPS, steps_for

PROJECT_DIR = os.path.dirname(os.path.abspath(__file__))
TASK_NAMES = {"watch": "CatalogPrep Watch", "daily": "CatalogPrep Daily"}


class Context:
    def __init__(self, run, cfg, options):
        self.run, self.cfg, self.options = run, cfg, options


def run_steps(ctx, names):
    """Run steps in order; stop at the first failure. Returns True when all succeed."""
    order = steps_for(ctx.run.status().get("mode", "manual"))
    for name in names:
        module, title = STEPS[name]
        pos = f"{order.index(name) + 1}/{len(order)}" if name in order else "-"
        heading(f"STEP {pos}: {title}  [{name}]")
        try:
            summary = module.run(ctx) or {}
        except StepError as e:
            ctx.run.step_failed(name, str(e))
            info(f"\nSTOPPED at step '{name}':\n{e}")
            return False
        except KeyboardInterrupt:
            ctx.run.step_failed(name, "interrupted")
            info("\nInterrupted.")
            return False
        except Exception as e:
            traceback.print_exc()
            ctx.run.step_failed(name, f"{type(e).__name__}: {e}")
            info(f"\nUnexpected error in step '{name}' (details above).")
            return False
        ctx.run.step_done(name, **summary)
    return True


def open_folder(path):
    if os.name == "nt" and sys.stdin is not None and sys.stdin.isatty():
        try:
            os.startfile(path)  # Windows Explorer
        except OSError:
            pass


# ---------------------------------------------------------------------------
# Mapping file + new run
# ---------------------------------------------------------------------------
def newest_mapping(cfg, given=None):
    if given:
        if not os.path.isfile(given):
            raise StepError(f"Mapping file not found: {given}")
        return given
    folder = cfg.path(cfg["mapping"]["folder"])
    found = find_mapping_files([folder])
    if not found:
        raise StepError(f"No mapping file in {folder}\n"
                        "Its name must contain 'image', 'link' and 'mapping' (or 'bakaul').")
    if len(found) > 1:
        info(f"{len(found)} mapping files found - using the newest.")
    return found[0]


def signature(path):
    st = os.stat(path)
    return f"{os.path.basename(path)}|{st.st_size}|{int(st.st_mtime)}"


def new_run(cfg, mapping, mode):
    session_date = date_from_filename(mapping)
    if not session_date:
        warn("No date in the mapping file name - using today's date.")
        session_date = datetime.now()
    run = Run.create(cfg["output_root"], session_date)
    shutil.copy2(mapping, run.folder("input"))
    run.set(mapping_file=os.path.basename(mapping), mapping_source=os.path.abspath(mapping),
            mapping_signature=signature(mapping), mode=mode, attempts=1)
    run.log(f"New {mode} run for mapping file: {os.path.basename(mapping)}")
    info(f"Run folder: {run.root}")
    return run


def todo_steps(run):
    return [s for s in steps_for(run.status().get("mode", "manual")) if not run.is_done(s)]


# ---------------------------------------------------------------------------
# Lock: never two runs at the same time
# ---------------------------------------------------------------------------
class Lock:
    def __init__(self, cfg, max_age_hours=6):
        os.makedirs(cfg["output_root"], exist_ok=True)
        self.path = os.path.join(cfg["output_root"], ".catprep.lock")
        self.max_age = max_age_hours * 3600

    def __enter__(self):
        if os.path.exists(self.path) and time.time() - os.path.getmtime(self.path) < self.max_age:
            with open(self.path, encoding="utf-8") as f:
                raise StepError(f"Another run is still working (started {f.read().strip()}). "
                                f"If that is wrong, delete {self.path}")
        with open(self.path, "w", encoding="utf-8") as f:
            f.write(datetime.now().strftime("%Y-%m-%d %H:%M"))
        return self

    def __exit__(self, *exc):
        try:
            os.remove(self.path)
        except OSError:
            pass


# ---------------------------------------------------------------------------
# Commands
# ---------------------------------------------------------------------------
def cmd_auto(args, cfg):
    auto_cfg = cfg.get("auto", {})
    if args.if_new and auto_cfg.get("headless_when_scheduled", True):
        cfg["export"]["headless"] = True
    with Lock(cfg):
        if args.resume:
            run = Run.open(cfg["output_root"], args.run)
        else:
            mapping = newest_mapping(cfg, args.mapping)
            run = None
            if args.if_new:
                sig = signature(mapping)
                previous = [r for r in Run.all_runs(cfg["output_root"])
                            if r.status().get("mapping_signature") == sig]
                if previous:
                    run = previous[0]
                    if run.status().get("mode") != "auto":
                        info(f"{os.path.basename(mapping)} is being handled by hand ({run.name}) - skipped.")
                        return True
                    if not todo_steps(run):
                        info(f"Nothing new: {os.path.basename(mapping)} was already done ({run.name}).")
                        return True
                    attempts = run.status().get("attempts", 1)
                    if attempts >= int(auto_cfg.get("max_attempts", 3)):
                        info(f"{run.name} failed {attempts} times - not retrying automatically. "
                             "Fix it and run: python catprep.py auto --resume")
                        return True
                    run.set(attempts=attempts + 1)
                    info(f"Retrying unfinished run {run.name} (attempt {attempts + 1}).")
            if run is None:
                run = new_run(cfg, mapping, "auto")
        if run.status().get("mode") != "auto":
            raise StepError(f"{run.name} is a manual (Claude) run - use prepare/finish for it.")
        ok = run_steps(Context(run, cfg, {"allow_missing": False}), todo_steps(run))
    report(cfg, run, ok)
    if ok:
        heading("AUTOMATIC RUN COMPLETE")
        info(f"Upload folder: {run.folder('upload')}")
        open_folder(run.folder("upload"))
    return ok


def report(cfg, run, ok):
    st = run.status()
    lines = [f"Run: {run.name}", f"Folder: {run.root}", f"Mapping file: {st.get('mapping_file')}", ""]
    for name in steps_for(st.get("mode", "manual")):
        s = st.get("steps", {}).get(name, {})
        extra = {k: v for k, v in s.items() if k not in ("state", "finished_at")}
        lines.append(f"{name:<8} {s.get('state', 'not run'):<8} " + ", ".join(f"{k}={v}" for k, v in extra.items()))
    subject = f"Catalogue Prep {'DONE' if ok else 'FAILED'}: {st.get('mapping_file')}"
    notify.send(cfg, subject, "\n".join(lines))


def cmd_prepare(args, cfg):
    with Lock(cfg):
        if args.resume:
            run = Run.open(cfg["output_root"], args.run)
            info(f"Resuming run: {run.root}")
            todo = [s for s in PREPARE if not run.is_done(s)]
            if not todo:
                info("All prepare steps are already done. Next: python catprep.py finish")
                return True
        else:
            run = new_run(cfg, newest_mapping(cfg, args.mapping), "manual")
            todo = PREPARE
        ok = run_steps(Context(run, cfg, {"manual_export": args.manual_export}), todo)
    if not ok:
        info("\nFix the problem, then continue with:  python catprep.py prepare --resume")
        return False
    heading("PREPARE COMPLETE")
    info(f"Files for Claude:   {run.folder('for_claude')}")
    info(f"Put cleaned files:  {run.folder('claude_done')}")
    info(r"Then run manual_claude\2_Finish.bat  (python catprep.py finish)")
    open_folder(run.folder("for_claude"))
    return True


def cmd_finish(args, cfg):
    with Lock(cfg):
        run = Run.open(cfg["output_root"], args.run)
        info(f"Run folder: {run.root}")
        not_ready = [s for s in PREPARE if not run.is_done(s)]
        if not_ready:
            raise StepError(f"Prepare is not finished for this run (missing: {not_ready}). "
                            "Run: python catprep.py prepare --resume")
        ok = run_steps(Context(run, cfg, {"allow_missing": args.allow_missing}), FINISH)
    if ok:
        heading("FINISH COMPLETE")
        info(f"Upload folder: {run.folder('upload')}")
        open_folder(run.folder("upload"))
    return ok


def cmd_step(args, cfg):
    with Lock(cfg):
        run = Run.open(cfg["output_root"], args.run)
        info(f"Run folder: {run.root}")
        ctx = Context(run, cfg, {"manual_export": args.manual_export, "allow_missing": args.allow_missing})
        return run_steps(ctx, [args.name])


def cmd_status(args, cfg):
    run = Run.open(cfg["output_root"], args.run)
    st = run.status()
    mode = st.get("mode", "manual")
    heading(f"RUN: {run.name}  ({'automatic' if mode == 'auto' else 'manual Claude'})")
    info(f"Folder:       {run.root}")
    info(f"Mapping file: {st.get('mapping_file', '-')}")
    info(f"Created:      {st.get('created_at', '-')}")
    info("")
    for n, name in enumerate(steps_for(mode), start=1):
        s = st.get("steps", {}).get(name, {})
        extra = {k: v for k, v in s.items() if k not in ("state", "finished_at")}
        details = ", ".join(f"{k}={v}" for k, v in extra.items())
        info(f"  {n}. {name:<7} {s.get('state', 'not started'):<12} {s.get('finished_at', ''):<20} {details}")
        if mode != "auto" and name == PREPARE[-1]:
            info("     --- Claude cleaning (manual) ---")
    return True


def cmd_runs(args, cfg):
    runs = Run.all_runs(cfg["output_root"])
    if not runs:
        info(f"No runs in {cfg['output_root']}")
    for r in runs:
        st = r.status()
        order = steps_for(st.get("mode", "manual"))
        done = sum(1 for s in order if st.get("steps", {}).get(s, {}).get("state") == "done")
        info(f"  {r.name:<40} {st.get('mode', 'manual'):<7} {done}/{len(order)} steps done")
    return True


def cmd_schedule(args, cfg):
    if os.name != "nt":
        raise StepError("Scheduling uses Windows Task Scheduler - run this on the Windows PC.")
    bat = os.path.join(PROJECT_DIR, "Auto_Scheduled.bat")
    auto_cfg = cfg.get("auto", {})

    def schtasks(*a):
        r = subprocess.run(["schtasks", *a], capture_output=True, text=True)
        info((r.stdout or r.stderr).strip())
        return r.returncode == 0

    if args.remove:
        for name in TASK_NAMES.values():
            schtasks("/Delete", "/TN", name, "/F")
        return True
    ok = True
    if args.watch:
        minutes = str(int(auto_cfg.get("watch_minutes", 10)))
        ok &= schtasks("/Create", "/TN", TASK_NAMES["watch"], "/SC", "MINUTE", "/MO", minutes,
                       "/TR", f'"{bat}"', "/F")
        info(f"Data folder is checked every {minutes} minutes.")
    if args.daily:
        at = args.at or auto_cfg.get("daily_time", "09:00")
        ok &= schtasks("/Create", "/TN", TASK_NAMES["daily"], "/SC", "DAILY", "/ST", at, "/TR", f'"{bat}"', "/F")
        info(f"Runs every day at {at} when there is a new mapping file.")
    if not (args.watch or args.daily):
        raise StepError("Choose --watch, --daily or --remove.")
    info("The task runs only while you are logged in to Windows. Output: scheduler.log in this folder.")
    return ok


def cmd_check(args, cfg):
    from openpyxl import load_workbook

    def ids_of(path):
        wb = load_workbook(path, read_only=True, data_only=True)
        ids = set()
        for ws in wb.worksheets:
            rows = ws.iter_rows(values_only=True)
            for _ in range(2):  # header can be in row 1 or 2
                header = next(rows, None) or []
                col = find_col(header, *PID_NAMES)
                if col is not None:
                    ids.update(p for p in (norm_pid(r[col]) for r in rows if col < len(r)) if p)
                    break
        wb.close()
        return ids

    for path in [args.file] + ([args.file2] if args.file2 else []):
        wb = load_workbook(path, read_only=True, data_only=True)
        heading(os.path.basename(path))
        for ws in wb.worksheets:
            first = [list(r) for r in ws.iter_rows(min_row=1, max_row=2, values_only=True)]
            info(f"Sheet '{ws.title}': {ws.max_row} rows x {ws.max_column} columns")
            for i, row in enumerate(first, start=1):
                info(f"  row {i}: {[v for v in row if v is not None][:12]}")
        wb.close()
        info(f"Unique Product IDs: {len(ids_of(path))}")
    if args.file2:
        a, b = ids_of(args.file), ids_of(args.file2)
        info(f"\nIn both: {len(a & b)} | only in first: {len(a - b)} | only in second: {len(b - a)}")
    return True


# ---------------------------------------------------------------------------
def main(argv=None):
    p = argparse.ArgumentParser(description="Catalogue Prep v2",
                                formatter_class=argparse.RawDescriptionHelpFormatter, epilog=__doc__)
    sub = p.add_subparsers(dest="command", required=True)

    sp = sub.add_parser("auto", help="fully automatic run")
    sp.add_argument("--if-new", action="store_true", help="only when there is a new mapping file")
    sp.add_argument("--resume", action="store_true", help="continue the newest (or --run) run")
    sp.add_argument("--mapping", help="use this mapping file instead of the newest in Data/")
    sp.add_argument("--run", help="run folder (only with --resume)")

    sp = sub.add_parser("schedule", help="Windows scheduled task")
    sp.add_argument("--watch", action="store_true", help="check Data every [auto] watch_minutes")
    sp.add_argument("--daily", action="store_true", help="every day at [auto] daily_time")
    sp.add_argument("--at", help="time for --daily, e.g. 09:00")
    sp.add_argument("--remove", action="store_true", help="remove the scheduled tasks")

    sp = sub.add_parser("prepare", help="manual Claude route: up to the Part files")
    sp.add_argument("--resume", action="store_true", help="continue the newest (or --run) run")
    sp.add_argument("--mapping", help="use this mapping file instead of the newest in Data/")
    sp.add_argument("--manual-export", action="store_true", help="put export files in yourself (no browser)")
    sp.add_argument("--run", help="run folder (only with --resume)")

    sp = sub.add_parser("finish", help="manual Claude route: Claude's files -> upload")
    sp.add_argument("--run", help="run folder (default: newest)")
    sp.add_argument("--allow-missing", action="store_true",
                    help="continue even if Claude's output is missing some rows")

    sp = sub.add_parser("step", help="re-run one step")
    sp.add_argument("name", choices=list(STEPS))
    sp.add_argument("--run", help="run folder (default: newest)")
    sp.add_argument("--manual-export", action="store_true")
    sp.add_argument("--allow-missing", action="store_true")

    for name in ("status", "runs"):
        sp = sub.add_parser(name)
        sp.add_argument("--run", help="run folder (default: newest)")

    sp = sub.add_parser("check", help="look inside an Excel file (read-only)")
    sp.add_argument("file")
    sp.add_argument("file2", nargs="?", help="second file: compare Product IDs")

    args = p.parse_args(argv)
    commands = {"auto": cmd_auto, "schedule": cmd_schedule, "prepare": cmd_prepare, "finish": cmd_finish,
                "step": cmd_step, "status": cmd_status, "runs": cmd_runs, "check": cmd_check}
    if args.command == "auto":
        info(f"\n===== {datetime.now():%Y-%m-%d %H:%M:%S}  catprep auto {' '.join(sys.argv[2:])}")
    try:
        cfg = config_mod.load(PROJECT_DIR)
        ok = commands[args.command](args, cfg)
    except StepError as e:
        info(f"\nERROR: {e}")
        ok = False
    except KeyboardInterrupt:
        info("\nInterrupted.")
        ok = False
    return 0 if ok else 1


if __name__ == "__main__":
    sys.exit(main())
